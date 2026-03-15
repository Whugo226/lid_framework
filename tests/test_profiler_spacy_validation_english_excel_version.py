import sys
import csv
import os
import math
import re
import numpy as np
import sklearn
import pytest
import spacy
import threading
from openpyxl import load_workbook
from collections import defaultdict
from scipy.stats import pearsonr, spearmanr

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src')))
from lid_toolkit.logic.profiler_spacy_bttm_up_approach import DeepProfiler

CSV_PATH = os.path.join(os.path.dirname(__file__), 'data', 'lingualyzer_ground_truth_english.xlsx')
RESULT_CSV_PATH = os.path.join(os.path.dirname(__file__), 'profiler_spacy_test_results_english_new4.csv')
CORRELATION_SUMMARY_CSV_PATH = os.path.join(os.path.dirname(__file__), 'profiler_spacy_correlation_english_summary.csv')
TOLERANCE = 1e-2

profiler = DeepProfiler()
nlp = profiler._make_pipeline('en')

test_results = []
results_lock = threading.Lock()

def get_profiler_stats(sentence):
    """Return a feature dict resembling a single-row profiling output.

    The ground-truth Excel file contains both "Doc" metrics and sentence-level
    overlap measures (Sent-Sent Avg).  The latter are calculated by
    ``DeepProfiler._calc_cross_level_overlaps`` when the profiler runs on
    an entire document, so we replicate that logic here.
    """
    doc = nlp(sentence)
    stats = profiler._extract_measures(doc, 'en')

    # paragraph measures mimic the behaviour of the full profiling
    # pipeline (see DeepProfiler._calc_paragraph_measures and the
    # document-level section in ``get_multilingual_profile``).
    paragraphs = re.split(r"\n\s*\n", sentence.strip())
    paragraph_count = len([p for p in paragraphs if p.strip()])
    stats['Paragraph count'] = paragraph_count
    if paragraph_count > 0:
        stats['Paragraph length'] = len(list(doc.sents)) / paragraph_count
    else:
        stats['Paragraph length'] = 0.0

    # compute sentence-based overlap metrics (adjacent pairs inside the text)
    sents_filtered = profiler._filter_valid_sentences(doc.sents)
    sent_docs = [sent.as_doc() for sent in sents_filtered]
    # we don't bother splitting into paragraphs for the test; an empty list is fine
    par_docs: list = []
    sent_to_par = profiler._build_sent_to_par_map(sentence, sents_filtered, par_docs)
    overlap_stats = profiler._calc_cross_level_overlaps(doc, par_docs, sent_docs, sent_to_par, 'en')

    stats.update(overlap_stats)
    return stats

def load_ground_truth():
    workbook = load_workbook(CSV_PATH, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            return

        normalized_headers = [str(h).strip() if h is not None else "" for h in headers]
        required = {"Measure", "Sentence", "Value"}
        missing = required.difference(normalized_headers)
        if missing:
            raise ValueError(f"Missing required columns in ground-truth workbook: {sorted(missing)}")

        header_to_index = {header: index for index, header in enumerate(normalized_headers)}
        for row in rows:
            measure_raw = row[header_to_index["Measure"]] if row else None
            sentence_raw = row[header_to_index["Sentence"]] if row else None
            value_raw = row[header_to_index["Value"]] if row else None

            measure = str(measure_raw).strip() if measure_raw is not None else ""
            sentence = str(sentence_raw).strip() if sentence_raw is not None else ""
            value = str(value_raw).strip() if value_raw is not None else ""
            if not measure or not sentence or not value:
                continue
            yield measure, sentence, value
    finally:
        workbook.close()

@pytest.mark.parametrize('measure,sentence,value', list(load_ground_truth()))
def test_profiler_measures_explicit(measure, sentence, value):
    status = "PASS"
    actual = None
    diff = None

    stats = get_profiler_stats(sentence)

    # The ground-truth CSV appends suffixes such as " (Doc)" or
    # " (Sent-Sent Avg)".  For document-level measures the profiler
    # returns keys without any suffix, but overlap metrics already include
    # the full label.  We try both forms when looking up the value in
    # *stats* so that the test doesn't skip the overlap measures.
    #
    # "measure_key_base" is the name with the trailing parenthesised
    # qualifier stripped (e.g. "Word overlap count (Sent-Sent Avg)" ->
    # "Word overlap count").  If that isn't found, we fall back to the
    # original string which may include the qualifier.
    measure_key_base = re.sub(r" \(.*\)$", "", measure)
    if measure_key_base in stats:
        measure_key = measure_key_base
    elif measure in stats:
        measure_key = measure
    else:
        # store an entry so the output file reflects missing keys
        with results_lock:
            test_results.append({
                "Measure": measure,
                "Sentence": sentence,
                "Status": "SKIP",
                "Expected": value,
                "Actual": None,
                "Difference": None
            })
        pytest.skip(f"Profiler output does not contain key '{measure_key_base}'")

    result = stats[measure_key]
    try:
        expected = float(value)
        actual = float(result)
        diff = abs(actual - expected)
        if measure.lower().endswith("count (doc)") or measure.lower().endswith("count"):
            if int(expected) != int(actual):
                status = "FAIL"
            assert int(expected) == int(actual)
        else:
            if diff >= TOLERANCE:
                status = "FAIL"
            assert diff < TOLERANCE
    except ValueError:
        expected = str(value)
        actual = str(result)
        diff = "N/A"
        if actual != expected:
            status = "FAIL"
        assert actual == expected
    finally:
        with results_lock:
            test_results.append({
                "Measure": measure,
                "Sentence": sentence,
                "Status": status,
                "Expected": expected,
                "Actual": actual,
                "Difference": diff
            })

def write_correlation_summary():
    if not test_results:
        return
    all_numeric_expected = []
    all_numeric_actual = []
    measure_data = defaultdict(lambda: {'expected': [], 'actual': []})
    for row in test_results:
        expected, actual = row['Expected'], row['Actual']
        if expected == 'N/A' or actual == 'N/A':
            continue
        try:
            e, a = float(expected), float(actual)
            all_numeric_expected.append(e)
            all_numeric_actual.append(a)
            measure_data[row['Measure']]['expected'].append(e)
            measure_data[row['Measure']]['actual'].append(a)
        except (TypeError, ValueError):
            continue

    with open(CORRELATION_SUMMARY_CSV_PATH, mode='w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=['Measure', 'MAE', 'Pearson', 'Spearman'])
        writer.writeheader()
        for measure, vals in sorted(measure_data.items()):
            expected_arr = np.array(vals['expected'])
            actual_arr = np.array(vals['actual'])
            mae = float(np.mean(np.abs(actual_arr - expected_arr)))
            if len(expected_arr) >= 2 and np.std(expected_arr) > 0 and np.std(actual_arr) > 0:
                pearson, _ = pearsonr(expected_arr, actual_arr)
                spearman, _ = spearmanr(expected_arr, actual_arr)
            else:
                pearson = spearman = float('nan')
            writer.writerow({'Measure': measure, 'MAE': mae, 'Pearson': pearson, 'Spearman': spearman})

    if len(all_numeric_expected) > 1:
        global_r, _ = pearsonr(all_numeric_expected, all_numeric_actual)
        print(f"\n{'='*40}\nPROFESSOR'S GLOBAL CORRELATION\nTotal: {len(all_numeric_expected)}\nr = {global_r:.4f}\n{'='*40}\n")

def write_results_to_csv():
    if not test_results:
        return
    with open(RESULT_CSV_PATH, mode="w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["Measure", "Sentence", "Status", "Expected", "Actual", "Difference"])
        writer.writeheader()
        writer.writerows(test_results)

@pytest.fixture(scope="session", autouse=True)
def write_results_fixture(request):
    yield
    write_results_to_csv()
    write_correlation_summary()